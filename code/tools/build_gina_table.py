#!/usr/bin/env python3
"""Build the round-1 review workbook for Gina (RelaxDB CPMG 10 proteins; AdK + CheY reviewed in depth).

Values come only from: the 10-protein pilot packet (anchored first pass), the AdK/CheY development review
(second-checked by Astra), the CheY Table 2 extraction, and the Stiller 2022 full text (PubMed Central).
"""
import csv
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = Path.home() / "dynamics_atlas_work/outputs/gina_review_round1"
OUT.mkdir(parents=True, exist_ok=True)

COLS = ["ID", "Review priority", "Protein (RelaxDB entry)", "Paper", "Link", "What we read", "Organism / construct vs entry",
        "Method", "Experiment and conditions", "Residues / region", "Dynamics observed", "Timescale (as reported)",
        "Population (as reported)", "Alternate-state structure", "Authors' conclusion", "Do not combine with / limits",
        "Location in paper", "Values checked", "Gina: correct? (Y/N)", "Gina: note"]

A2 = "second reader checked numbers against paper"
A1 = "first pass, anchored quote; not second-checked"

ROWS = [
# ---------------- AdK (Aquifex aeolicus) ----------------
["AQADK-1", "A - please check", "AQADK_CPMG (Aquifex AdK, O66490)", "Henzler-Wildman et al. 2007, Nature 450:838",
 "https://doi.org/10.1038/nature06410", "Main text pp.838-842 + SI Figs S6/S7", "Target (Aquifex WT)", "Crystallography",
 "Apo crystals; three molecules in the asymmetric unit", "ATP lid, AMP lid, hinges", "Conformers at different positions along the lid-closing path",
 "-", "-", "Apo conformers along the closing trajectory (this paper's crystal structures)",
 "Closed-like conformations are geometrically accessible without ligand", "Crystal contacts stabilize these conformers; not solution populations or rates",
 "pp.838-839, Fig.2", A2, "", ""],
["AQADK-2", "A - please check", "AQADK_CPMG", "Henzler-Wildman et al. 2007, Nature 450:838", "https://doi.org/10.1038/nature06410",
 "Main text + SI Fig S6", "Target (Aquifex WT)", "NMR relaxation dispersion (15N CPMG)",
 "Apo; deuterated sample 10 C; protonated sample 20 C", "Exchanging residues near hinges/lids (Fig.3 map)",
 "us-ms exchange", "kex = 950 +/- 150 s-1 (10 C, deuterated); 20 C: only a compatible range 2000-7000 s-1", "Not determined for apo",
 "-", "Intrinsic lid/hinge motion in apo enzyme", "CPMG alone does not identify a fully closed endpoint; 20 C value is a range",
 "pp.839-840; SI Fig.S6", A2, "", ""],
["AQADK-3", "A - please check", "AQADK_CPMG", "Henzler-Wildman et al. 2007, Nature 450:838", "https://doi.org/10.1038/nature06410",
 "Main text p.840", "Target (Aquifex WT)", "NMR relaxation dispersion (15N CPMG)", "Mg-Ap5A saturated; 20 C; protonated",
 "Lids/hinges", "Opening/closing exchange in inhibitor-bound enzyme", "kex = 300 +/- 150 s-1", "Closed state 95 +/- 2% (two-site fit)",
 "-", "Bound enzyme favors the closed form", "Bound-state values; not apo population, not turnover rate",
 "p.840, Fig.3b", A2, "", ""],
["AQADK-4", "A - please check", "AQADK_CPMG", "Henzler-Wildman et al. 2007, Nature 450:838", "https://doi.org/10.1038/nature06410",
 "Main text p.841 + methods", "Target species, dye-labelled Y52C/V145C construct (35% of WT activity)", "smFRET",
 "Apo; room temperature", "Lid distance (labels at 52/145)", "Lid opening/closing transitions",
 "Opening 6500 +/- 500 s-1; closing 2000 +/- 200 s-1 (corrected)", "Not reported; our two-state estimate 23.5% closed (2000/8500)",
 "-", "Closed-like states are sampled on us-ms scale", "Labelled construct and model-derived population; do not compare directly with NMR 95%",
 "p.841, Fig.5", A2, "", ""],
["AQADK-5", "A - please check", "AQADK_CPMG", "Henzler-Wildman et al. 2007, Nature 450:838", "https://doi.org/10.1038/nature06410",
 "Main text p.842", "Target species, Y52C spin-labelled", "PRE", "Apo", "Lid region near label",
 "Short-distance contributions consistent with closed-like sampling", "-", "Authors: PRE does not give populations",
 "-", "Supports closed-like sampling spatially", "Distance weighting; do not read as high closed occupancy", "p.842, Fig.3e-g", A2, "", ""],
["AQADK-6", "A - please check", "AQADK_CPMG", "Henzler-Wildman et al. 2007, Nature 450:913", "https://doi.org/10.1038/nature06407",
 "Abstract, captions, author-posted excerpts, SI Table S3", "Aquifex (thermophile) compared with E. coli (mesophile)", "NMR spin relaxation (ps-ns) + MD",
 "Temperature comparison between homologs", "Hinge regions (SI Table S3 uses E. coli-equivalent numbering)",
 "Fast hinge fluctuations linked to slower lid motion", "ps-ns (order parameters)", "-", "-",
 "Fast local motion facilitates large-scale lid motion", "Order parameters are not us-ms populations; residue numbers need mapping to Aquifex",
 "Abstract; SI p.9 Table S3 footnote a", A2, "", ""],
["AQADK-7", "A - please check", "AQADK_CPMG", "Pelz et al. 2016, Nat Commun 7:10848", "https://doi.org/10.1038/ncomms10848",
 "Results + methods (archived full text)", "Target species, Cys insertions for optical-tweezers handles", "Single-molecule force spectroscopy",
 "Optical tweezers ~10 pN; ATP/ADP/Ap5A, AP4A, AP6A", "Lid domains", "Partly closed substrate-bound state; ligand-specific closing",
 "Substrate kinetics too fast to resolve", "Zero-force population extrapolated, not measured",
 "ADP/ATP: dominant minimum with both lids half closed; AP5A/AP4A full closing; AP6A mainly partial", "Ligand identity changes the preferred closed state",
 "Under load; do not substitute for unloaded apo WT values", "Results: nucleotide binding; Fig.3c/d, Fig.4", A2, "", ""],
["AQADK-8", "A - please check", "AQADK_CPMG", "Wolf-Watz et al. 2004, Nat Struct Mol Biol 11:945", "https://doi.org/10.1038/nsmb821",
 "Full text (PDF supplied by Alex)", "Target (Aquifex 'thermoAdk')", "NMR relaxation dispersion (15N CPMG) during catalysis",
 "Substrate-saturated, turning over; 20 C; 500/600 and 800 MHz", "Both lids (residues in both lids fit one kex)",
 "Lid opening/closing during turnover; lid opening is rate-limiting",
 "kex = 1615 +/- 100 s-1; kopen = 44 +/- 20 s-1; kclose = 1571 +/- 100 s-1; kcat = 30 +/- 10 s-1 (20 C)", "Not extracted (closing much faster than opening; equilibrium far toward closed)",
 "-", "Slower lid opening alone explains the 9-fold lower activity at 20 C", "50 C turnover kex > 10000 s-1 (data not shown); E. coli values are row AQADK-8b",
 "Results: 'Determination of lid-closing and opening rates'", "read in this round (Claude)", "", ""],
["AQADK-8b", "A - please check", "AQADK_CPMG (homolog evidence)", "Wolf-Watz et al. 2004, Nat Struct Mol Biol 11:945", "https://doi.org/10.1038/nsmb821",
 "Full text (PDF supplied by Alex)", "HOMOLOG: E. coli AdK ('mesoAdk')", "NMR relaxation dispersion (15N CPMG) during catalysis",
 "Substrate-saturated, turning over; 20 C", "Both lids", "Lid opening/closing during turnover",
 "kex = 1660 +/- 90 s-1; kopen = 286 +/- 85 s-1; kclose = 1374 +/- 110 s-1; kcat = 263 +/- 30 s-1", "-", "-",
 "kopen matches kcat in both enzymes", "Homolog: do not attach to Aquifex", "Results: 'Determination of lid-closing and opening rates'", "read in this round (Claude)", "", ""],
["AQADK-9", "A - please check", "AQADK_CPMG (homolog evidence)", "Stiller et al. 2022, Nature 603:528", "https://doi.org/10.1038/s41586-022-04468-9",
 "Full text (PubMed Central)", "HOMOLOG: Geobacillus stearothermophilus AdK (inferred from closed reference PDB 4QBH; Zn/Co lid site) - not Aquifex",
 "PCS-CPMG (paramagnetic NMR relaxation dispersion)", "Turnover conditions: 20 mM Mg-ADP, pH 7.0, 25 C; Co/Zn forms",
 "AMP lid opens ~15 deg (~50%); ATP lid moves ~1.8 A", "High-energy state during catalysis",
 "kex = 1428 +/- 83 s-1; lid opening 180 +/- 36 s-1", "High-energy state 12.6 +/- 2.5%",
 "Solved high-energy structure: partially closed (RMSD 2.67 A to closed vs 7.03 A to open)",
 "Small lid openings, primed for substrate binding/product release", "Homolog: do not attach these numbers to Aquifex",
 "Results; Methods (NMR)", "read in this round; species inferred", "", ""],
# ---------------- CheY ----------------
["CHEY-1", "A - please check", "CHEY_CPMG (E. coli CheY, P0AE67)", "McDonald, Boyer & Lee 2012, Structure 20:1363",
 "https://pmc.ncbi.nlm.nih.gov/articles/PMC3552614/", "Results, methods, Table 2", "Target (WT, unphosphorylated)",
 "NMR relaxation dispersion (15N CPMG)", "No Mg; 1 mM EDTA; pH 7.0; 15 C; 600/700 MHz",
 "R18 R19 A36 D38 N62 D64 E67 L68 L69 T87 N94 V107 L127 (13 residues; see sheet CheY_Table2)",
 "Local exchange, not a single concerted switch", "kex 96-2440 s-1 (local fits)", "pI 0.67-0.99 (local fits; no global fraction)",
 "-", "Segmental motions, not a two-state concerted switch", "Do not mix with the 10 mM Mg table (Table 1); D13 too broadened to fit",
 "P15, Table 2", A2, "", ""],
["CHEY-2", "A - please check", "CHEY_CPMG", "McDonald, Boyer & Lee 2012, Structure 20:1363", "https://pmc.ncbi.nlm.nih.gov/articles/PMC3552614/",
 "Results P10-P12", "Target (WT, unphosphorylated)", "NMR relaxation dispersion (15N CPMG)", "10 mM Mg; 15 C",
 "11 fitted residues (Table 1)", "Exchange persists; group fits do not support one switch", "Table 1 (not extracted here)", "Table 1 (not extracted here)",
 "Tested intermediate 2B1J did not give expected shift correlation", "Simple concerted switch rejected", "Mg occupancy differs from EDTA sample",
 "P10-P12, P35-P39, Table 1", A2, "", ""],
["CHEY-3", "A - please check", "CHEY_CPMG", "McDonald et al. 2013", "https://pmc.ncbi.nlm.nih.gov/articles/PMC3703654/",
 "Results P8-P25, methods", "Target WT vs BeFx-bound (phosphorylation mimic)", "NMR relaxation dispersion + fast backbone/methyl relaxation",
 "10 mM Mg; pH 7.0; CPMG 15 C; fast relaxation 25 C",
 "Active site (Rex down); beta4-alpha4 loop and near Y106 (residual); FliM interface (new low Rex); T87/A88 methyls more rigid",
 "Activation redistributes slow and fast motions", "-", "-", "-", "Fast and slow dynamics colocalize",
 "Unactivated sample not Mg-saturated at 10 mM, BeFx sample is; baseline CPMG reused from 2012; BeFx is a mimic",
 "P8-P17, P26-P31", A2, "", ""],
# ---------------- other 8 proteins (first pass) ----------------
["BLAC-1", "2 - overview", "BLAC_CPMG (M. tuberculosis BlaC, P9WKD3)", "BlaC dynamics study 2020, AAC", "https://doi.org/10.1128/AAC.02025-19",
 "Full text", "Target", "NMR relaxation dispersion (15N CPMG)", "Free mature BlaC; pH 6.4; 298 K; 600/850 MHz",
 "Active site: S70, S130, T235, T237 (Ambler numbering)", "us-ms exchange in active site", "kex = 860 +/- 60 s-1 (two-state group fit)",
 "Not accurately determined", "-", "Active-site dynamics increase upon inhibitor binding", "Do not pool with inhibitor conditions; Ambler numbering needs mapping",
 "Results, Fig.3-4", A1, "", ""],
["BLAC-2", "2 - overview", "BLAC_CPMG", "BlaC dynamics study 2020, AAC", "https://doi.org/10.1128/AAC.02025-19", "Full text", "Target",
 "Relaxation dispersion + crystal comparison", "Clavulanate adduct or avibactam", "Omega loop; K73", "Inhibitor-dependent exchange",
 "-", "-", "Ligand-dependent crystal conformations (candidate)", "Structural explanation proposed", "Peak splitting and sample change complicate clavulanate data",
 "Discussion, Fig.6-7", A1, "", ""],
["BLVRB-1", "2 - overview", "BLVRB_CPMG (human BLVRB, P30043)", "Biliverdin reductase B dynamics 2018, JMB", "https://doi.org/10.1016/j.jmb.2018.06.015",
 "Full text", "Target", "NMR relaxation dispersion (CPMG)", "Apo vs NADP+-bound; 20 C", "R78 loop/active site incl. S111, T164; bound: D36/S37/R39",
 "Coenzyme binding redistributes exchange", "Fast-exchange fits", "Not separable from shift differences", "-",
 "Dynamics coupled to coenzyme binding", "Same residue exchanging in two states is not the same process", "Results P14, Fig.4", A1, "", ""],
["BLVRB-2", "2 - overview", "BLVRB_CPMG", "BLVRB catalytic cycle 2023, Front Mol Biosci", "https://doi.org/10.3389/fmolb.2023.1244587",
 "Full text + Table 2 PDF", "Target", "NMR relaxation dispersion, temperature-global fit", "900 MHz; 0/5/10 C (apo also 20 C)",
 "Residue group II", "Group-specific exchange", "Group II 10 C kex ~457 s-1 (apo), 3884 +/- 450 s-1 (NADP+)",
 "Group II 96.9 +/- 0.4% (apo), 96.4 +/- 3.1% (NADPH), 99.3 +/- 0.1% (NADP+)", "Apo cryo and room-temperature crystals (R78 side chain)",
 "Partial correspondence of exchange shifts with apo/holo differences", "Model estimates; apo/NADPH ranges overlap", "Table 2 (PDF p.11), Fig.2", A1, "", ""],
["CYPA-0", "2 - overview", "CYPA_CPMG (human CypA, P62937)", "Eisenmesser et al. 2005, Nature 438:117", "https://doi.org/10.1038/nature04105",
 "Full text (PDF supplied by Alex)", "Target WT; mutants R55A, R55K, F113W, H70A, K82A separate", "NMR relaxation dispersion (15N and 13C CPMG)",
 "Free enzyme vs during turnover (substrate present); temperature not extracted in this round",
 "Group I probes incl. L39, R55, I57, L98, S99, A103, G110, E120; group II loop 65-84 (D66, F67, T68, N71)",
 "Free enzyme already samples the motion seen during catalysis",
 "Free: group I kex = 1140 +/- 200 s-1; loop 65-84 kex = 2260 +/- 200 s-1. Turnover: kex = 2730 +/- 220 s-1", "Minor state upper limit about 7% (free WT)",
 "-", "Intrinsic dynamics of the enzyme underlie catalysis", "Mutants shift populations at similar kex; turnover kex compared with cis/trans isomerization 2500 +/- 500 s-1",
 "Main text, Fig.1-3; Methods", "read in this round (Claude)", "", ""],
["CYPA-1", "2 - overview", "CYPA_CPMG (human CypA, P62937)", "Fraser et al. 2009, Nature 462:669", "https://doi.org/10.1038/nature08615",
 "Full text", "Target WT and S99T", "Room-temperature crystallography + NMR", "WT cryo vs room temperature; S99T",
 "R55, M61, L98, S99, F113 network", "Alternate side-chain rotamers in the network", "us-ms (WT); S99T slows group I", "-",
 "Room-temperature alternate rotamers (minor state candidates)", "Hidden alternative structures essential for catalysis", "Original 2005 CPMG source paper not read in full",
 "Results P5-P14, Fig.1-4", A1, "", ""],
["CYPA-2", "2 - overview", "CYPA_CPMG", "CypA temperature-jump X-ray scattering 2019", "https://pmc.ncbi.nlm.nih.gov/articles/PMC6815256/",
 "Full text", "Target WT; S99T and D66N/R69H controls", "Time-resolved X-ray scattering (T-jump)", "~10.7 C jump to 26.7 C",
 "Core and 65-85 loop", "Two distinct relaxation processes", "k1 = (9.8 +/- 1.2) x 10^5 s-1; k2 = (2.5 +/- 0.2) x 10^4 s-1", "-", "-",
 "Core and loop motions are distinct", "Global relaxation constants; not equilibrium NMR kex", "Results P14-P20, Fig.5-6", A1, "", ""],
["VHR-1", "2 - overview", "VHR_CPMG (human DUSP3, P51452)", "VHR variable insert study 2020, Biochemistry", "https://doi.org/10.1021/acs.biochem.0c00245",
 "Full text", "Target WT", "NMR relaxation dispersion (15N CPMG)", "Apo; pH 6.5; 25 C; 600/700/800 MHz",
 "VI: A63, R66, N72; acid loop: K89, D92, E95, F96, L98, S99", "Coupled VI/acid-loop exchange", "kex = 1530 +/- 200 s-1 (combined fit)",
 "Major 98.8 +/- 0.1%", "No empty-pocket structure in this paper; 2025 study adds apo crystals 8TK6, 9DJ9",
 "Coupled allosteric motions", "RelaxDB source is a dissertation table; equivalence not confirmed", "Results P23-P24, Fig.3", A1, "", ""],
["ARGKIN-1", "2 - overview", "ARGKIN_CPMG (Limulus arginine kinase, P51541)", "Arginine kinase motions 2011, Biochemistry", "https://doi.org/10.1021/bi101664u",
 "Full text", "Target (3 positions differ from UniProt)", "NMR relaxation dispersion, multi-temperature", "Substrate-free; 15-30 C",
 "Loop L8 (I182-G209); NTD/CTD interface; hinge", "Domain and loop motions", "L8 kex ~500-1100 s-1; interface ~600-1400 s-1",
 "L8 minor 3-6%", "Open and transition-state-analogue crystals (references)", "Rate-limiting domain and loop motions",
 "Closing-path link is an assumed correspondence", "Results P21-P24, Table 1", A1, "", ""],
["ARGKIN-2", "2 - overview", "ARGKIN_CPMG", "Arginine kinase intrinsic dynamics 2009", "https://pmc.ncbi.nlm.nih.gov/articles/PMC2826323/",
 "Full text", "Target", "NMR relaxation dispersion (residue and group fits)", "Substrate-free; pH 6.5; 25 C",
 "35 exchanging residues; hinge D88, G90, G92", "Hinge/interface/loop exchange", "Hinge kex 1930 +/- 350 s-1; interface 800 +/- 100 s-1",
 "Hinge 0.4 +/- 0.1%; interface 1.1 +/- 0.1%", "-", "Motions commensurate with turnover", "Group/model-specific values", "Results; Table", A1, "", ""],
["RNASE-1", "2 - overview", "RNASE_CPMG (bovine RNase A)", "Beach et al. 2005, JACS 127:9167", "https://doi.org/10.1021/ja0514949",
 "Original PDF read remotely", "Target", "NMR relaxation dispersion", "Apo vs pTppAp; pH 6.3; 283-298 K; 500/600 MHz",
 "Loop 1 (14-24), loop 4 (64-71)", "Conserved us-ms motion in apo and substrate-mimic states",
 "293 K: apo 1390 +/- 70 s-1; complex 1880 +/- 200 s-1", "Major 90-93% (apo), 90-96% (complex) over 283-298 K",
 "7RSA/1U1B as structural references", "Apo samples bound-like conformations", "Condition-specific; ranges", "p.9172-9176, Fig.4, Fig.8", A1, "", ""],
["RNASE-2", "2 - overview", "RNASE_CPMG", "What's in your buffer? 2013, Biochemistry", "https://doi.org/10.1021/bi400973e",
 "Full text", "Target", "NMR relaxation dispersion vs buffer anion", "298 K; 600 MHz; MES vs 100 mM phosphate / 80 mM sulfate / 100 mM acetate",
 "Probe-dependent (e.g. A4, A19)", "Solute changes the detected exchange", "-", "-", "-",
 "Buffer solutes alter ms motions", "This is the phosphate confound Gina mentioned", "Results P97-P113, Fig.3-6", A1, "", ""],
["RNASE-3", "2 - overview", "RNASE_CPMG", "RNase A energy landscape 2015, J Phys Chem B", "https://doi.org/10.1021/acs.jpcb.5b00212",
 "Full text", "Target", "CPMG and R1rho vs pH", "pH 4.5-7.0; 10-25 C", "Cluster 1 and cluster 2",
 "Cluster 1: no CPMG dispersion at pH 4.5 but R1rho detects exchange", "-", "Not determined at pH 5.5/4.5", "-",
 "Protonation-coupled exchange proposed", "Flat CPMG is not 'no exchange'", "Results P105-P135, Fig.2/5", A1, "", ""],
["HTPDZ-1", "2 - overview", "HTPDZ_CPMG (HtrA2 PDZ, O43464 352-458)", "HtrA2 substrate recognition 2024, Nat Commun", "https://doi.org/10.1038/s41467-024-48997-5",
 "Full text", "Target isolated PDZ (not full-length)", "Backbone and methyl relaxation dispersion", "500 uM; 283 K / 298 K",
 "N-terminus, alpha5 + linker to beta15, alpha7 + linker; beta17-beta18 loop (temperature-dependent)", "ms exchange",
 "tau_ex 0.98 ms (283 K), 0.42 ms (298 K); methyl 1.19 ms (298 K)", "Minor 16.1% (283 K), 15.2% (298 K); methyl 16.6 +/- 0.4%",
 "-", "PDZ exchange quantified", "Isolated-PDZ values; not the full-length activation state", "Results Par7-Par11, Fig.1-2", A1, "", ""],
["HTPDZ-2", "2 - overview", "HTPDZ_CPMG", "HtrA2 substrate processing 2022", "https://pmc.ncbi.nlm.nih.gov/articles/PMC9170070/",
 "Full text", "Full-length S306A/I441V (not the entry construct)", "State-sensitive methyl NMR", "40 C; substrate titration",
 "Protease-PDZ interface", "Closed / open-inactive / open-active states", "-", "Active population varies with substrate", "-",
 "PDZ engagement opens but does not activate", "Different construct and conditions", "Results, Fig.2", A1, "", ""],
["KRAS-1", "2 - overview", "KRAS_CPMG (K-Ras4B 1-169)", "K-Ras excited state 2023, NSMB", "https://doi.org/10.1038/s41594-023-01070-z",
 "Full text", "Target WT (G12C/G12D separate)", "CPMG + CEST", "GTP-bound; 298 K", "P-loop; Switch I 30-38; Switch II 60-76",
 "GDP-like, partly disordered excited state", "kex = 400 +/- 12 s-1 (WT)", "Excited state 1: 10.15 +/- 0.47%", "-",
 "Excited state resembles GDP-like switch disorder", "Not chemical conversion to GDP; RelaxDB row ligand not documented", "Results, Table 1, Fig.3-5", A1, "", ""],
["KRAS-2", "2 - overview", "KRAS_CPMG", "Mg-free state in K-Ras 2022", "https://pmc.ncbi.nlm.nih.gov/articles/PMC9804424/",
 "Full text", "Target WT 1-169", "CPMG + Mg-release reference shifts", "GDP/Mg; 25 C; pH 7.4", "Switch regions",
 "Mg-release-like minor state", "kex = 164 +/- 28 s-1", "pE = 2.8 +/- 0.2%", "-", "Mg2+-free state matters for nucleotide exchange",
 "GDP-state minor state; not the 2023 GTP state 1", "Table 1", A1, "", ""],
]

CHEY_NOTE = "CheY Table 2 (McDonald et al. 2012): unphosphorylated WT, no Mg, 1 mM EDTA, pH 7.0, 15 C, 600/700 MHz. Local two-state fits per residue."

METHOD_CLASSES = ["NMR relaxation dispersion", "NMR spin relaxation (ps-ns)", "Exchange broadening / missing peaks",
                  "RDC / PRE / PCS ensemble NMR", "Multi-temperature / ensemble crystallography", "Time-resolved structural",
                  "FRET / smFRET", "EPR / DEER", "HDX-MS", "MD simulation", "PDB structure comparison (apo/holo)"]
COVER = {  # found in rows of this workbook; everything else = not searched yet in round 1
    "AQADK_CPMG": {0: "AQADK-2,3,8; AQADK-8b, AQADK-9 (homolog)", 1: "AQADK-6", 3: "AQADK-5 (PRE); AQADK-9 (PCS, homolog)", 4: "AQADK-1", 6: "AQADK-4", 9: "AQADK-6", 10: "AQADK-1"},
    "BLAC_CPMG": {0: "BLAC-1,2", 10: "BLAC-2"},
    "BLVRB_CPMG": {0: "BLVRB-1,2", 4: "BLVRB-2"},
    "CYPA_CPMG": {0: "CYPA-0, CYPA-1", 4: "CYPA-1", 5: "CYPA-2"},
    "VHR_CPMG": {0: "VHR-1", 9: "VHR-1 (MD in paper)", 10: "VHR-1 (2025 apo crystals)"},
    "ARGKIN_CPMG": {0: "ARGKIN-1,2", 3: "2011 RDC study (not in rows)", 10: "ARGKIN-1"},
    "RNASE_CPMG": {0: "RNASE-1,2,3", 10: "RNASE-1 (7RSA/1U1B)"},
    "CHEY_CPMG": {0: "CHEY-1,2,3", 1: "CHEY-3", 10: "CHEY-2 (2B1J)"},
    "HTPDZ_CPMG": {0: "HTPDZ-1", 1: "HTPDZ-2 (methyl)"},
    "KRAS_CPMG": {0: "KRAS-1,2", 1: "KRAS-1"},
}

ISSUES = [
    ["AQADK_CPMG", "Naming", "GitHub JSON calls the entry MesoAdK; sequence equals Aquifex aeolicus O66490", "Which organism/sample is this entry?"],
    ["BLAC_CPMG", "Label", "A Zenodo figure caption calls BLAC 'TEM-1'; UniProt P9WKD3 is M. tuberculosis BlaC (265/266 residues match 43-307)", "Is the caption a typo?"],
    ["RNASE_CPMG", "Organism", "Entry uses P61824 (UniProt now: Bison bison); sequence also equals mature bovine RNase A (P61823); paper studies bovine", "Switch accession to P61823?"],
    ["ARGKIN_CPMG", "Sequence", "Entry differs from current UniProt P51541 at positions 103, 112, 116", "Construct variant or UniProt revision?"],
    ["KRAS_CPMG", "Identity", "No UniProt id; sequence = K-Ras4B (P01116-2) residues 1-169; ligand/temperature of the row not documented", "Which nucleotide state and paper?"],
    ["HTPDZ_CPMG", "Construct", "Entry 109 aa = UniProt 352-458 + 'LE'; paper construct described as 349-458; BMRB 51320 entity is 156 aa", "Which construct produced the labels?"],
    ["VHR_CPMG", "Source", "Labels from a dissertation table (7.10.3); equivalence with the 2020 paper not confirmed", "Is the thesis table the same dataset?"],
    ["CPMG labels", "Version drift", "Hugging Face relaxdb_cpmg labels differ from metadata xlsx (09dec2024): AQADK 19, BLAC 19, BLVRB 7, CYPA 6, VHR 26, RNASE 4, KRAS 51 positions; ARGKIN, CHEY, HTPDZ identical", "Which version is current?"],
    ["Phosphate column", "Conflict", "xlsx 'Phosphate buffer?': 24 Y, 3 N, 106 blank. Hugging Face 'phosphate': True for 9 entries (3 of the 24 Y, plus 6 blanks); False for 21 of the 24 Y", "What does each field mean?"],
]

SEND = [
    ["AQADK_CPMG", "Henzler-Wildman 2007a", "https://doi.org/10.1038/nature06410", "read", "key paper"],
    ["AQADK_CPMG", "Henzler-Wildman 2007b", "https://doi.org/10.1038/nature06407", "partial", "key paper"],
    ["AQADK_CPMG", "Wolf-Watz 2004", "https://doi.org/10.1038/nsmb821", "read (PDF)", "key paper"],
    ["AQADK_CPMG", "Pelz 2016", "https://doi.org/10.1038/ncomms10848", "read", "supporting"],
    ["AQADK_CPMG", "Stiller 2022 (homolog)", "https://doi.org/10.1038/s41586-022-04468-9", "read", "homolog evidence; check organism"],
    ["AQADK_CPMG", "Kerns 2015", "https://pmc.ncbi.nlm.nih.gov/articles/PMC4318763/", "main text E. coli; Aquifex SI not read", "check SI"],
    ["CHEY_CPMG", "McDonald 2012", "https://pmc.ncbi.nlm.nih.gov/articles/PMC3552614/", "read (SI not read)", "key paper"],
    ["CHEY_CPMG", "McDonald 2013", "https://pmc.ncbi.nlm.nih.gov/articles/PMC3703654/", "read (SI not read)", "key paper"],
    ["CHEY_CPMG", "2020 FliM priming study", "-", "not read (fetch failed)", "lead"],
    ["CYPA_CPMG", "Eisenmesser 2005 (original CPMG)", "https://doi.org/10.1038/nature04105", "read (PDF)", "key paper"],
]


CASE_ENTRIES = ["BLAC_CPMG", "BLVRB_CPMG", "CYPA_CPMG", "VHR_CPMG", "ARGKIN_CPMG", "RNASE_CPMG", "HTPDZ_CPMG", "KRAS_CPMG",
                "CHEY_CPMG", "AQADK_CPMG"]
PREFIX = {"BLAC_CPMG": "BLAC", "BLVRB_CPMG": "BLVRB", "CYPA_CPMG": "CYPA", "VHR_CPMG": "VHR", "ARGKIN_CPMG": "ARGKIN",
          "RNASE_CPMG": "RNASE", "HTPDZ_CPMG": "HTPDZ", "KRAS_CPMG": "KRAS", "CHEY_CPMG": "CHEY", "AQADK_CPMG": "AQADK"}
CASES = Path.home() / "dynamics_atlas_work/cases"


HAND_DOI = {"CHEY-1": "10.1016/j.str.2012.05.008", "CHEY-2": "10.1016/j.str.2012.05.008", "CHEY-3": "10.1016/j.jmb.2013.04.029"}
HAND_KEEP = set()
CHECKED_AUTO = "automatic check passed: quoted sentence and numbers found in the paper text; not second-read by a person"
WORK_LEDGER = Path.home() / "dynamics_atlas_work/ledger"
# Exported row ids are fixed once assigned (2026-10-17): judgments cite them, so excluding or adding rows later
# must never renumber other rows.
REGISTRY_PATH = WORK_LEDGER / "row_id_registry.json"
REGISTRY = json.loads(REGISTRY_PATH.read_text()) if REGISTRY_PATH.exists() else {}


def best_case(entry):
    """Latest layered (or Claude) case for entry whose integrity check passed."""
    cands = [p for p in CASES.glob(f"R_r*_{entry}_*")] + [p for p in CASES.glob(f"L_r*_{entry}_*")] + [p for p in CASES.glob(f"C_{entry.split('_')[0]}_claude")] + \
            [p for p in CASES.glob(f"p1r1_{entry}")]
    good = []
    for p in cands:
        try:
            sc = json.loads((p / "scorecard.json").read_text())
        except Exception:
            continue
        ov = json.loads((WORK_LEDGER / "integrity_overrides.json").read_text()).get(p.name, []) \
            if (WORK_LEDGER / "integrity_overrides.json").exists() else []
        errs = sc.get("integrity", {}).get("errors") or []
        overridden = bool(errs) and all(any(o["row"] == e.get("row") and o["kind"] == e.get("kind") for o in ov) for e in errs)
        if (sc.get("integrity", {}).get("pass") or overridden) and (p / "rows.jsonl").exists():
            good.append(p)
    return sorted(good, key=lambda p: p.stat().st_mtime)[-1] if good else None


A_METHODS = {"nmr_relaxation_dispersion", "nmr_spin_relaxation", "nmr_exchange_broadening"}
B_METHODS = {"pdb_structure_comparison", "hdx_ms", "fret", "epr_deer", "time_resolved_structural", "xray_multitemp_ensemble", "nmr_ensemble_restraints"}


DENATURED = re.compile(r"denatur|unfolded state|\burea\b|guanidin", re.I)


def tier(r):
    if DENATURED.search(" ".join(str(r.get(k) or "") for k in ("title", "construct_or_mutation", "experiment", "other_conditions"))):
        return "C"
    near = r.get("species_relation") in ("target", "same_protein_other_construct")
    full = r.get("access") != "abstract_only"
    if near and full and r.get("method_class") in A_METHODS:
        return "A"
    if near and full and r.get("method_class") in B_METHODS:
        return "B"
    return "C"


INTERNAL = re.compile(r"\.json|\bXML\b|bioc|visually (checked|inspected|verified)|cached|worksheet saved|re-extracted|candidates\.json|"
                      r"retrieval failed|correction DOI|in artefact_note|original XML|source XML", re.I)


def plain(text):
    """Drop internal bookkeeping (file names, retrieval notes) from a Gina-visible cell."""
    text = str(text or "")
    text = text.replace("仅图示", "figure only")  # fill workers sometimes write this Chinese note into English fields
    text = re.sub(r"(,|;)\s*visually (checked|inspected|verified)", "", text)
    text = re.sub(r"visually (checked|inspected|verified)", "checked against the figure", text)
    text = re.sub(r"\s*\([^()]*(\.json|XML|cached)[^()]*\)", "", text)
    text = re.sub(r"the supplied entry( sequence)?", "the RelaxDB entry sequence", text)
    text = re.sub(r"~?/?[\w./-]*(sources_cache|dynamics_atlas_work)/[\w./-]+", "", text)
    text = re.sub(r"\bcached\s+", "", text)
    text = re.sub(r";?\s*± signs lost in text extraction[^);]*", "", text)
    text = re.sub(r",\s*;", ";", text)
    text = re.sub(r"\(\s*[,;]?\s*\)", "", text).replace(" ,", ",").replace("( ", "(")
    text = text.replace("homolog of the supplied Bison bison accession", "RelaxDB accession P61824 is Bison bison; mature sequence identical")
    parts = [x for x in re.split(r";\s*", text) if x.strip() and not INTERNAL.search(x)]
    return "; ".join(parts).strip()


MERGED = {}


def rows_from_case(case, entry, extra=()):
    out = []
    pre = PREFIX[entry]
    raw = []
    tov_all = json.loads((WORK_LEDGER / "tier_overrides.json").read_text()) if (WORK_LEDGER / "tier_overrides.json").exists() else {}
    corr_all = json.loads((WORK_LEDGER / "row_corrections.json").read_text()) if (WORK_LEDGER / "row_corrections.json").exists() else {}
    exc_all = json.loads((WORK_LEDGER / "row_exclusions.json").read_text()) if (WORK_LEDGER / "row_exclusions.json").exists() else {}
    for c in [case, *extra]:  # main case plus hand-curated supplements share the tier limits
        for r in (json.loads(x) for x in (c / "rows.jsonl").read_text().splitlines() if x.strip()):
            if r.get("row_id") in exc_all.get(c.name, {}):
                continue
            r["_case"], r["_force_c"] = c.name, tov_all.get(c.name, {}).get(r.get("row_id"))
            for k, v in corr_all.get(c.name, {}).get(r.get("row_id"), {}).items():  # source-verified corrections
                if k.endswith("_append"):
                    f = {"caveat_append": "caveat_note"}.get(k, k[:-7])
                    r[f] = f"{r.get(f) or ''}; {v}".strip("; ")
                else:
                    r[k] = v
            raw.append(r)
    idd = json.loads((WORK_LEDGER / "identity_decisions.json").read_text()).get(entry, {}) \
        if (WORK_LEDGER / "identity_decisions.json").exists() else {}
    for r in raw:
        if idd and r.get("species_relation") == "homolog":
            r["species_relation"] = idd["homolog_to"]
            r["caveat_note"] = ((r.get("caveat_note") or "") + " " + idd["note"]).strip()
    # rows of one paper that repeat the same values (e.g. joint CPMG + CEST fit split in two) become one row
    merged, key_of = [], {}
    for r in raw:
        k = (r.get("doi"), r.get("construct_or_mutation"), r.get("ligand_or_partner"), r.get("temperature"), r.get("other_conditions"),
             r.get("residues_author"), r.get("timescale"), r.get("population"))
        if (r.get("timescale") or r.get("population")) and k in key_of:
            m = key_of[k]
            MERGED[(r["_case"], r.get("row_id"))] = m.get("row_id")
            m["experiment"] = f"{m.get('experiment') or ''} + {r.get('experiment') or ''}"
            continue
        key_of[k] = r
        merged.append(r)
    raw = merged

    def base_tier(r):
        if r["_force_c"]:
            return "C"
        t = tier(r)
        if t == "A" and not (r.get("residues_author") or r.get("timescale") or r.get("population")):
            return "B"  # nothing Gina can check in an A row
        return t

    def capped(rows, total, per_paper):
        # one paper should not fill a tier: keep at most per_paper rows per DOI, then the total cap
        seen, keep = {}, []
        for r in rows:
            d = r.get("doi") or r.get("title")
            if seen.get(d, 0) < per_paper:
                seen[d] = seen.get(d, 0) + 1
                keep.append(r)
        return {id(r) for r in keep[:total]}
    a = [r for r in raw if base_tier(r) == "A"]
    a.sort(key=lambda r: (r.get("species_relation") != "target", not (r.get("timescale") or r.get("population"))))
    keep_a = capped(a, 15, 2)
    b = [r for r in raw if base_tier(r) == "B" or (base_tier(r) == "A" and id(r) not in keep_a)]
    b.sort(key=lambda r: not (r.get("timescale") or r.get("population")))
    keep_b = capped(b, 20, 3)
    label = {"A": "A - please check", "B": "B - supporting", "C": "C - appendix"}
    used_ids = set()
    for r in raw:
        t = base_tier(r)
        if t == "A" and id(r) not in keep_a:
            t = "B"
        if t == "B" and id(r) not in keep_b:
            t = "C"
        ph = str(r.get("ph") or "").strip()
        ph = ph if ph.lower().startswith("ph") or not ph else "pH " + ph
        cond = "; ".join(x for x in [r.get("temperature"), ph, r.get("buffer"), r.get("ligand_or_partner"), r.get("other_conditions")] if x)
        link = ("https://doi.org/" + r["doi"]) if r.get("doi") else ""
        key = f"{r['_case']}|{r.get('row_id')}"
        rid = REGISTRY.get(key)
        if rid is None:  # first export of this row: plain id unless taken; then the id is fixed for good
            taken = used_ids | set(REGISTRY.values())
            rid = f"{pre}-{r.get('row_id')}"
            if rid in taken:  # a supplement that reuses r1, r2 ... gets its own prefix
                rid = f"{pre}-{r['_case'].split('_', 2)[-1]}-{r.get('row_id')}"
            REGISTRY[key] = rid
            REGISTRY_PATH.write_text(json.dumps(REGISTRY, indent=0, ensure_ascii=False))
        used_ids.add(rid)
        limits = " ".join(x for x in [r.get("artefact_note"), r.get("caveat_note"), r.get("caveat"), r["_force_c"]] if x)
        out.append([rid, label[t], entry, r.get("title") or "", link, r.get("access") or "",
                    plain(f"{r.get('species_relation')}; {r.get('construct_or_mutation') or ''}"), r.get("method_class") or "",
                    plain(f"{r.get('experiment') or ''}; {cond}"), r.get("residues_author") or "", r.get("dynamics_type") or "",
                    r.get("timescale") or "", r.get("population") or "", r.get("alt_state_structure") or "",
                    r.get("author_claim") or "", plain(limits), plain(r.get("locator")) or "see paper",
                    CHECKED_AUTO, "", ""])
    return out


def case_coverage(case, exported=None):
    try:
        cov = json.loads((case / "coverage.json").read_text()).get("classes", {})
    except Exception:
        return None
    keys = ["nmr_relaxation_dispersion", "nmr_spin_relaxation", "nmr_exchange_broadening", "nmr_ensemble_restraints",
            "xray_multitemp_ensemble", "time_resolved_structural", "fret", "epr_deer", "hdx_ms", "md_simulation", "pdb_structure_comparison"]
    out = []
    for k in keys:
        v = cov.get(k, {})
        if v.get("status") == "found":
            pre = PREFIX.get((json.loads((case / "input.json").read_text()).get("entry_id")), "")
            ids = []
            for rid in v.get("row_ids", []):
                rid = MERGED.get((case.name, rid), rid)  # rows merged at export point to the kept row
                full = f"{pre}-{rid}"
                if (exported is None or full in exported) and full not in ids:
                    ids.append(full)
            out.append("found: " + ", ".join(ids) if ids else "found in search; rows excluded or merged")
        elif v.get("status") == "searched_not_found":
            out.append("searched (Europe PMC method queries); no relevant paper selected")
        else:
            out.append("not searched yet")
    return out


def sheet(ws, header, rows, widths):
    ws.append(header)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F766E")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"


JUDG = Path.home() / "dynamics_atlas_work/outputs/judgment_v0"


def ranges(v):
    """[30,31,32,40] -> '30-32, 40'; strings pass through."""
    if not isinstance(v, list):
        return str(v or "")
    nums = sorted({int(x) for x in v if str(x).lstrip("-").isdigit()})
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(str(nums[i]) if i == j else f"{nums[i]}-{nums[j]}")
        i = j + 1
    return ", ".join(out)


ALT_APO = Path.home() / "dynamics_atlas_work/outputs/alt_apo_states_v0.json"
ALT_RULE = ("Our reading of the 9-17 definition (alternate APO state = an APO state with slow motion, slow enough for "
            "separate NMR peaks); please correct it. yes = in the reference state a minor state is measured with a population "
            "and a rate (CPMG/CEST/R1rho) or a second set of peaks is seen; possible = exchange detected but no population, or a "
            "population only in another condition; no_evidence = measured, not detected in a stated window; unknown = not "
            "measured or search insufficient.")


def alt_apo_sheet(wb):
    """One row per protein: is there an alternate APO state, by Gina's 9-17 definition (outputs/alt_apo_states_v0.json)."""
    if not ALT_APO.exists():
        return
    rows = []
    for a in json.loads(ALT_APO.read_text()):
        sm = "; ".join(f"{m.get('pdb')}: {m.get('what') or m.get('what the authors say it resembles') or ''} ({m.get('evidence_row', '')})"
                       for m in a.get("structure_match", []) or [])
        rows.append([a.get("entry"), a.get("call"), a.get("literal_two_peak", ""), a.get("basis"), a.get("reference_state"), a.get("kex"), a.get("pB"),
                     a.get("regime"), a.get("residues"), sm, a.get("fallback_annotation"),
                     "; ".join(map(str, a.get("other_conditions", []) or [])), a.get("caveats"),
                     ", ".join(a.get("evidence_rows", []) or []), ", ".join(a.get("judgments", []) or []), "", ""])
    rows = [[plain(c) if isinstance(c, str) else c for c in r] for r in rows]
    sheet(wb.create_sheet("Alternate_APO_states", 1), ["Protein", "Alternate APO state? (our rule)", "Two separate peaks? (literal 9-17 definition)", "Basis", "Reference state", "kex", "pB", "Exchange regime", "Residues",
               "Structure match (PDB)", "Fallback annotation", "Other conditions", "Caveats", "Evidence rows", "Judgments",
               "Gina: correct? (Y/N)", "Gina: note"], rows, [14, 12, 30, 60, 40, 20, 14, 24, 36, 40, 50, 40, 40, 30, 24, 12, 30])


def judgment_sheets(wb):
    """Per-protein judgment pages from outputs/judgment_v0/<ENTRY>.json (only files whose check passed)."""
    files = []
    for f in sorted(JUDG.glob("*_CPMG.json")):
        chk = JUDG / f"{f.stem}.check.json"
        reviewed = json.loads((WORK_LEDGER / "judgment_reviewed.json").read_text()) if (WORK_LEDGER / "judgment_reviewed.json").exists() else {}
        if chk.exists() and json.loads(chk.read_text()).get("pass") and f.stem in reviewed:  # only science-reviewed proteins
            files.append(f)
    if not files:
        return
    summ, rows = [], []
    for f in files:
        d = json.loads(f.read_text())
        ps = d.get("protein_summary", {})
        pw = ps.get("per_window", {})
        summ.append([f.stem, plain(ps.get("wt_apo_answer", "")), pw.get("us-ms", ""), pw.get("ps-ns", ""), pw.get("slower_than_ms", ""),
                     pw.get("undetermined", ""), ps.get("buffer_note", ""), "; ".join(map(str, ps.get("open_conflicts", []) or [])),
                     "; ".join(map(str, ps.get("questions_for_gina", []) or [])), "", ""])
        for j in d.get("judgments", []):
            st, det, reg = j.get("sample_state", {}), j.get("detection", {}), j.get("region", {})
            rows.append([j.get("judgment_id"), f.stem, j.get("time_window"), f"{st.get('construct', '')}; {st.get('binding_state', '')}",
                         st.get("conditions", ""), ranges(reg.get("residues_entry")) or reg.get("residues_author", ""),
                         j.get("direction"), j.get("verdict"), det.get("method", ""),
                         "; ".join(x for x in [det.get("kex") and f"kex {det.get('kex')}", det.get("pB") and f"pB {det.get('pB')}"] if x),
                         det.get("detection_window", ""), j.get("reason", ""), j.get("evidence_strength", ""),
                         ", ".join(j.get("evidence_rows", []) or []), ", ".join(j.get("conflicts_with", []) or []), "", ""])
    summ = [[plain(c) if isinstance(c, str) else c for c in r] for r in summ]
    rows = [[plain(c) if isinstance(c, str) else c for c in r] for r in rows]
    sheet(wb.create_sheet("Protein_summary", 1),
          ["Protein", "Reference state (WT apo; K-Ras: WT, nucleotide-bound, no partner): is it dynamic?", "us-ms window", "ps-ns window", "slower than ms", "window undetermined",
           "Buffer / condition notes", "Open conflicts", "Questions for Gina", "Gina: correct? (Y/N)", "Gina: note"],
          summ, [14, 60, 50, 40, 30, 30, 40, 40, 40, 12, 30])
    sheet(wb.create_sheet("Judgments", 2),
          ["ID", "Protein", "Time window", "Construct; binding state", "Conditions", "Residues (entry numbering)", "Direction",
           "Verdict", "Method", "Rates / population", "Detection window (if not detected)", "Reason", "Evidence strength",
           "Evidence rows", "Conflicts with", "Gina: correct? (Y/N)", "Gina: note"],
          rows, [11, 13, 13, 26, 30, 22, 18, 18, 26, 24, 30, 50, 18, 22, 14, 12, 30])


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    for line in [
        "RelaxDB CPMG literature evidence - round 1 draft (2026-10-11)",
        "",
        "Please spend ~15 minutes on one or two proteins of your choice, rows marked 'A - please check'.",
        "We would like to know three things:",
        "1. Which key papers are missing?",
        "2. Which numbers, residues or sample conditions are attached to the wrong experiment?",
        "3. Are the RelaxDB data issues (sheet 'RelaxDB_issues') real?",
        "",
        "Mark column 'Gina: correct? (Y/N)' and add a note where useful.",
        "One row = one paper x one sample state x one experiment. Every number is copied as reported in the paper.",
        "'A - please check': NMR dynamics evidence on the target protein or the same protein in another construct (at most 15 per protein, at most 2 per paper; more rows of the same paper move to B).",
        "'B - supporting': structural and other-method evidence, plus extra NMR rows beyond the A limit (at most 20 per protein).",
        "'C' rows are in sheet 'Appendix_rows': MD, homologs, abstract-only papers, denatured-state studies and overflow.",
        "NEW (10-20): sheet 'Alternate_APO_states' gives one call per protein for the question set on 9-17. " + ALT_RULE,
        "NEW: sheets 'Protein_summary' and 'Judgments' answer 'is it dynamic, where, on what timescale' for all 10 proteins, split by time window and sample state. 'Not detected' means not seen in that detection window, not 'no dynamics'. Every judgment cites the evidence rows it rests on.",
        "Sheet 'Method_coverage' shows, per method, which rows were found or what was searched; the last column lists known search gaps.",
        "Sheet 'Papers_to_send' lists the key papers per protein and the papers we could not read in full (paywalled, or supplement not read).",
        "Rows were extracted by an automated reader (OpenAI Codex) or by Claude, and checked automatically: the quoted sentence and every number must appear in the source text. Every row was then re-read against its paper by two independent models (Codex and Claude; Claude only for BlaC and CypA) to catch values attached to the wrong sample; each flagged row was checked against the paper and corrected where wrong. No person has second-read the rows yet.",
    ]:
        ws.append([line])
    ws.column_dimensions["A"].width = 120
    ws["A1"].font = Font(bold=True, size=14)

    rows_all, used = [], {}
    for e in CASE_ENTRIES:
        c = best_case(e)
        if c:
            used[e] = c
    case_dois = {}
    for e, c in used.items():
        case_dois[e] = {json.loads(x).get("doi", "").lower() for x in (c / "rows.jsonl").read_text().splitlines() if x.strip()} - {""}
    for r in ROWS:
        ent = r[2].split(" ")[0]
        if ent in used and r[0] not in HAND_KEEP and (str(r[1]).startswith("2") or HAND_DOI.get(r[0], "-") in case_dois[ent]
                                                      or any(d in str(r[4]).lower() for d in case_dois[ent])):
            continue  # overview rows and papers already in the case are replaced by the case rows
        rows_all.append(r)
    for e, c in used.items():
        sups = []
        for sup in sorted(CASES.glob(f"S_{PREFIX[e]}_*")):
            try:
                sc = json.loads((sup / "scorecard.json").read_text())
            except Exception:
                continue
            row_errors = [x for x in sc.get("integrity", {}).get("errors", []) if not str(x.get("kind", "")).startswith("coverage")]
            if (sup / "rows.jsonl").exists() and not row_errors:  # single-paper supplements: coverage errors are expected
                sups.append(sup)
        rows_all.extend(rows_from_case(c, e, sups))
    order = {"1": 0, "A": 1, "2": 2, "B": 3}
    main = [r for r in rows_all if not str(r[1]).startswith("C")]
    main.sort(key=lambda r: (r[2], order.get(str(r[1])[0], 4)))
    appendix = [r for r in rows_all if str(r[1]).startswith("C")]
    sheet(wb.create_sheet("Evidence_rows"), COLS, main, [9, 14, 24, 30, 30, 22, 28, 24, 32, 34, 30, 32, 30, 32, 32, 36, 22, 20, 12, 30])
    sheet(wb.create_sheet("Appendix_rows"), COLS, appendix, [9, 14, 24, 30, 30, 22, 28, 24, 32, 34, 30, 32, 30, 32, 32, 36, 22, 20, 12, 30])
    (OUT / "case_sources.json").write_text(json.dumps({e: c.name for e, c in used.items()}, indent=1))

    rows = list(csv.DictReader(open(Path.home() / "dynamics_atlas_work/outputs/development_review_20261005/CHEY_TABLE2.csv")))
    ws2 = wb.create_sheet("CheY_Table2")
    ws2.append([CHEY_NOTE])
    ws2.append([])
    hdr = ["Residue", "kex (s-1)", "delta omega (ppm)", "pI", "R2 600 MHz (s-1)", "R2 700 MHz (s-1)", "chi2", "Location", "Gina: correct? (Y/N)"]
    ws2.append(hdr)
    for r in rows:
        ws2.append([r["entry_residue"] + r["author_number"], r["kex_s_inv"], r["delta_omega_ppm"], r["pI"], r["R2_600_s_inv"], r["R2_700_s_inv"], r["chi_squared"], "Table 2", ""])
    for c in ws2[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F766E")
    for i, w in enumerate([10, 14, 16, 12, 16, 16, 8, 10, 16], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    cov = []
    for p, found in COVER.items():
        cc = case_coverage(used[p], {r[0] for r in rows_all}) if p in used else None
        gaps = json.loads((WORK_LEDGER / "search_gaps.json").read_text()) if (WORK_LEDGER / "search_gaps.json").exists() else {}
        cov.append([p] + (cc if cc else [found.get(i, "not searched yet") for i in range(len(METHOD_CLASSES))]) + [gaps.get(p, "")])
    sheet(wb.create_sheet("Method_coverage"), ["Protein"] + METHOD_CLASSES + ["Known search gaps"], cov, [16] + [20] * len(METHOD_CLASSES) + [60])
    sheet(wb.create_sheet("RelaxDB_issues"), ["Entry", "Type", "What we found", "Question for Gina", "Gina: answer"], ISSUES, [16, 14, 80, 36, 30])
    send = list(SEND)
    for e, c in used.items():
        try:
            for x in json.loads((c / "not_read_in_full.json").read_text()):
                send.append([e, x.get("title") or "", ("https://doi.org/" + x["doi"]) if x.get("doi") else "", "abstract only", x.get("reason") or "need full text"])
        except Exception:
            pass
    sheet(wb.create_sheet("Papers_to_send"), ["Protein", "Paper", "Link", "Our access", "Role"], send, [16, 30, 50, 32, 30])

    xlsx = OUT / "RelaxDB_CPMG_evidence_round1.xlsx"
    judgment_sheets(wb)
    alt_apo_sheet(wb)
    wb.save(xlsx)
    with open(OUT / "evidence_rows.csv", "w", newline="") as f:  # every exported row: main sheet then appendix
        w = csv.writer(f)
        w.writerow(COLS)
        w.writerows(main + appendix)
    print(xlsx, len(main), "main rows,", len(appendix), "appendix rows")


if __name__ == "__main__":
    main()
