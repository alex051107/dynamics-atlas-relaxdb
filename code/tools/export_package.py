#!/usr/bin/env python3
"""Export evaluated development cases; source checks are not scientific acceptance."""
import argparse
import csv
from datetime import datetime, timezone
from html import escape
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
from urllib.parse import urlsplit
import xml.etree.ElementTree as ET

WORK = Path(__file__).resolve().parents[1]
EVAL = Path.home() / 'dynamics_atlas_eval'
RUNTIME = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3'
ISSUES = [
    ('all', 'Residue annotations differ between dataset versions for 7/10 CPMG proteins.'),
    ('all', 'The phosphate-buffer column contradicts between two metadata tables.'),
    ('AQADK_CPMG', 'MesoAdK versus Aquifex identity.'),
    ('BLAC_CPMG', 'BLAC is labelled TEM-1.'),
    ('RNASE_CPMG', 'RNase species needs confirmation.'),
    ('ARGKIN_CPMG', 'Three residue types differ from UniProt.'),
    ('KRAS_CPMG', 'UniProt accession is missing.'),
    ('HTPDZ_CPMG', 'HtrA2 construct needs confirmation.'),
]
REPLY = ('We have prepared the AdK and CheY literature evidence and marked unresolved sample and condition questions.\n'
         'Which key studies have we missed?\n'
         'Which quantities or experimental conditions need correction?\n'
         'Could you confirm whether the reported dataset identity, residue annotation and buffer discrepancies are real?\n'
         'The highlighted papers are included for your review, with the remaining evidence in the appendix.\n')


def read(path):
    return json.loads(path.read_text())


def rows(path):
    return [json.loads(line) for line in path.read_text().splitlines() if line.strip()]


def validate_case(case, dev_ids, cache):
    card, run, score = [read(case / n) for n in ('input.json', 'run.json', 'scorecard.json')]
    entry = card['entry_id']
    if entry not in dev_ids or score.get('visibility') != 'full':
        raise ValueError('Only evaluated cases within the current export scope allowed')
    if run.get('status') != 'done' or score.get('integrity', {}).get('pass') is not True:
        raise ValueError('Completed worker and passing integrity required')
    evidence, coverage = rows(case / 'rows.jsonl'), read(case / 'coverage.json')['classes']
    methods = read(WORK / 'config/recipe.json')['methods']
    if set(coverage) != set(methods) or not evidence:
        raise ValueError('Nonempty evidence and all eleven coverage classes required')
    ids = [r['row_id'] for r in evidence]
    if len(ids) != len(set(ids)):
        raise ValueError('Duplicate evidence row id')
    for row in evidence:
        source = (cache / row['source_file']).resolve()
        if cache.resolve() not in source.parents or not source.is_file():
            raise ValueError('Source is absent or outside the source cache')
        if row.get('access') == 'abstract_only' and any(re.search(r'\d', str(row.get(k) or ''))
                for k in ('temperature', 'ph', 'timescale', 'population')):
            raise ValueError('Abstract-only evidence contains a numerical claim')
    for method, cell in coverage.items():
        if cell.get('status') == 'found' and (not cell.get('row_ids') or any(
                not any(r['row_id'] == rid and r['method_class'] == method for r in evidence)
                for rid in cell['row_ids'])):
            raise ValueError('Coverage does not bind to matching evidence rows')
    table2 = []
    if entry == 'CHEY_CPMG':
        # Fresh case provenance + explicit types; Gina still checks scientific extraction.
        started = run.get('started', '').split('_')[0]
        if not re.fullmatch(r'\d{14}', started) or started < '20261002000000':
            raise ValueError('CheY requires a newly run case under the current order')
        lead = re.search(r'https?://\S+', card.get('source_lead', ''))
        identity = re.sub(r'[^a-z0-9]', '', urlsplit(lead.group()).path.split('/')[-1].lower()) if lead else ''
        if not identity: raise ValueError('CheY primary source identity missing')
        for r in evidence:
            if not re.search(r'\btable\s*2\b', r.get('locator', ''), re.I):
                continue
            if r.get('species_relation') != 'target' or r.get('access') != 'full_text':
                continue
            source = cache / r['source_file']; xml = source if source.suffix == '.xml' else source.with_suffix('.xml')
            if not xml.is_file(): raise ValueError('CheY Table 2 requires the primary XML for verification')
            tree = ET.fromstring(xml.read_text())
            article = tree if tree.tag == 'article' else tree.find('.//article')
            if article is None: raise ValueError('Primary XML has no article object')
            paper_ids = [re.sub(r'[^a-z0-9]', '', ''.join(e.itertext()).lower())
                         for e in article.findall('./front/article-meta/article-id')]
            receipt_path = xml.with_suffix('.source.json')
            aliases = read(receipt_path).get('identifier_aliases', {}) if receipt_path.exists() else {}
            resolved = [re.sub(r'[^a-z0-9]', '', record['doi'].lower()) for key, record in aliases.items()
                        if re.sub(r'[^a-z0-9]', '', key.lower()) == identity and record.get('authority_url')]
            tables = [e for e in article.iter('table-wrap') if re.search(r'\btable\s*2\b',
                      ' '.join(e.find('label').itertext()) if e.find('label') is not None else '', re.I)]
            if not (identity in paper_ids or any(i in paper_ids for i in resolved)) or not tables:
                raise ValueError('CheY source does not match the supplied primary paper/Table 2')
            typed = re.findall(r'\b([ACDEFGHIKLMNPQRSTVWY])(\d+)\b', r.get('residues_author', ''))
            positions = re.findall(r'\b\d+\b', str(r.get('residues_entry', '')))
            if not typed or len(typed) != len(positions) or '-' in r.get('residues_author', ''):
                raise ValueError('CheY Table 2 needs individual typed and mapped residues')
            for (letter, number), position in zip(typed, positions):
                position = int(position)
                if not 1 <= position <= len(card['sequence']) or card['sequence'][position - 1] != letter:
                    raise ValueError('CheY Table 2 residue type mismatch')
                table2.append({'row_id': r['row_id'], 'author': letter + number,
                               'entry_position': position, 'source_file': r['source_file'], 'locator': r['locator']})
        if not table2:
            raise ValueError('CheY primary Table 2 extraction/type check is missing')
    summary = (case / 'summary.md').read_text()
    if len(summary.split()) > 600 or any(rid not in ids for rid in re.findall(r'\[(r\d+)\]', summary)):
        raise ValueError('Summary length or row references invalid')
    return {'entry': entry, 'card': card, 'run': run, 'case': str(case), 'rows': evidence,
            'coverage': coverage, 'summary': summary, 'table2_type_checks': table2,
            'not_read': read(case / 'not_read_in_full.json')}


def paper_list(evidence):
    papers = {}
    for r in evidence:
        key = r.get('doi') or r.get('pmid') or r['title']
        p = papers.setdefault(key, {'title': r['title'], 'doi': r.get('doi'), 'pmid': r.get('pmid'), 'row_ids': []})
        p['row_ids'].append(r['row_id'])
    # ponytail: use worker order for review priorities; no automatic scientific ranking.
    return list(papers.values())


def write_xlsx(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    book = Workbook(); book.remove(book.active)
    for name, records in sheets.items():
        ws = book.create_sheet(name)
        columns = list(dict.fromkeys(k for r in records for k in r)) or ['entry_id', 'status']
        ws.append(columns)
        for record in records:
            ws.append([json.dumps(record.get(k), ensure_ascii=False) if isinstance(record.get(k), (dict, list))
                       else record.get(k) for k in columns])
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str): cell.data_type = 's'  # source text cannot become an Excel formula
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in ws[1]: cell.font = Font(bold=True)
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 30
    book.save(path)


def render(output, cases):
    output.mkdir(parents=True, exist_ok=True)
    evidence = [r for c in cases for r in c['rows']]
    coverage = [{'entry_id': c['entry'], 'method_class': method, **cell}
                for c in cases for method, cell in c['coverage'].items()]
    not_read = [{'entry_id': c['entry'], **p} for c in cases for p in c['not_read']]
    issues = [{'entry_id': entry, 'issue': text, 'status': 'USER-REPORTED, UNVERIFIED',
               'source': 'User-adopted execution order; requires collaborator confirmation'} for entry, text in ISSUES]
    write_xlsx(output / 'atlas.xlsx', {'evidence_rows': evidence, 'method_coverage': coverage,
                                     'relaxdb_issues': issues, 'not_read_in_full': not_read})
    columns = list(dict.fromkeys(k for r in evidence for k in r))
    with (output / 'atlas_rows.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=columns); writer.writeheader(); writer.writerows(evidence)
    review = {}
    for c in cases:
        papers = paper_list(c['rows']); review[c['entry']] = {'priority': papers[:8], 'appendix': papers[8:]}
        body = escape(c['summary'])
        body = re.sub(r'\[(r\d+)\]', r'<a href="#\1">[\1]</a>', body).replace('\n\n', '</p><p>').replace('\n', '<br>')
        table = ''.join('<tr><td>' + escape(method.replace('_', ' ')) + '</td><td>' + escape(cell['status'])
                        + '</td><td>' + escape(', '.join(cell.get('row_ids', []))) + '</td></tr>'
                        for method, cell in c['coverage'].items())
        references = ''.join('<li>' + escape(p['title']) + (' — doi:' + escape(p['doi']) if p['doi'] else '') + '</li>'
                             for p in papers[:8])
        appendix = ''.join('<section id="' + escape(r['row_id']) + '"><h3>' + escape(r['row_id'])
                          + '</h3><pre>' + escape(json.dumps(r, ensure_ascii=False, indent=2)) + '</pre></section>' for r in c['rows'])
        html = ('<!doctype html><meta charset="utf-8"><title>' + escape(c['entry']) + '</title><style>'
                'body{font:14px/1.35 system-ui;max-width:950px;margin:24px auto}table{border-collapse:collapse;width:100%}'
                'td{border-bottom:1px solid #ddd;padding:3px}pre{white-space:pre-wrap;overflow-wrap:anywhere}'
                '@page{size:A4;margin:12mm}@media print{body{font:9pt/1.15;margin:0}.appendix{display:none}'
                '.summary{columns:2;column-gap:7mm}li,table{font-size:8pt}a{color:inherit;text-decoration:none}}</style>'
                '<h1>' + escape(c['card'].get('name') or c['entry']) + '</h1><div class="summary"><p>' + body + '</p></div>'
                '<h2>Method coverage</h2><table>' + table + '</table><h2>Papers for review</h2><ol>' + references
                + '</ol><p>Experimental evidence remains specific to the reported sample and detection window.</p>'
                '<div class="appendix"><h2>Evidence rows and source locators</h2>' + appendix + '</div>')
        (output / (c['entry'] + '.html')).write_text(html)
    (output / 'review_papers.json').write_text(json.dumps(review, ensure_ascii=False, indent=2))
    if {c['entry'] for c in cases} == {'AQADK_CPMG', 'CHEY_CPMG'}:
        (output / 'REPLY_GINA_EN.md').write_text(REPLY)
    manifest = {'entries': [c['entry'] for c in cases], 'integrity_pass': True,
                'built_at': datetime.now(timezone.utc).isoformat(), 'scientific_acceptance': 'pending_Gina',
                'source_cases': [{'case': c['case'], 'version': c['run']['prompt_version']} for c in cases],
                'chey_table2_type_checks': [r for c in cases for r in c['table2_type_checks']]}
    (output / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2))


def self_test():
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory(dir=WORK / 'ledger/tests') as temp:
        root = Path(temp); case = root / 'case'; cache = root / 'cache'
        case.mkdir(); cache.mkdir(); (cache / 'fixture.txt').write_text('Synthetic protocol fixture, not scientific evidence.')
        (cache / 'fixture.xml').write_text('<article><front><article-meta><article-id>fixture-paper</article-id></article-meta></front>'
                                         '<body><table-wrap><label>Table 2</label><table>A1 C2</table></table-wrap></body></article>')
        evidence = [{'row_id': 'r1', 'title': '=FORMULA <script>', 'source_file': 'fixture.txt', 'method_class': 'fret',
                     'access': 'full_text', 'locator': 'Table 2', 'species_relation': 'target',
                     'residues_author': 'A1, C2', 'residues_entry': '1, 2'}]
        objects = {'input.json': {'entry_id': 'CHEY_CPMG', 'name': '<script>CheY', 'sequence': 'AC',
                                 'source_lead': 'https://example.org/fixture-paper#tbl2'},
                   'run.json': {'status': 'done', 'started': '20261002120000_a', 'prompt_version': 'fixture'},
                   'scorecard.json': {'visibility': 'full', 'integrity': {'pass': True}},
                   'coverage.json': {'classes': {m: {'status': 'searched_not_found', 'queries': []}
                                      for m in read(WORK / 'config/recipe.json')['methods']}}, 'not_read_in_full.json': []}
        objects['coverage.json']['classes']['fret'] = {'status': 'found', 'row_ids': ['r1']}
        for n, v in objects.items(): (case / n).write_text(json.dumps(v))
        (case / 'rows.jsonl').write_text(json.dumps(evidence[0]) + '\n'); (case / 'summary.md').write_text('Fixture [r1].')
        c = validate_case(case, {'CHEY_CPMG'}, cache); render(root / 'out', [c])
        book = load_workbook(root / 'out/atlas.xlsx')
        assert book.sheetnames == ['evidence_rows', 'method_coverage', 'relaxdb_issues', 'not_read_in_full']
        assert next(cell for row in book['evidence_rows'] for cell in row if cell.value == '=FORMULA <script>').data_type == 's'
        assert '<script>' not in (root / 'out/CHEY_CPMG.html').read_text()
        assert all(r['status'] == 'USER-REPORTED, UNVERIFIED' for r in [dict(zip(
            [c.value for c in book['relaxdb_issues'][1]], row)) for row in book['relaxdb_issues'].iter_rows(min_row=2, values_only=True)])
        original = (cache/'fixture.xml').read_text()
        (cache/'fixture.xml').write_text('<pmc-articleset>' + original.replace('<article-id>fixture-paper</article-id>',
                                        '<article-id pub-id-type="doi">10.fixture</article-id>') + '</pmc-articleset>')
        receipt = {'identifier_aliases': {'fixture-paper': {'doi':'10.fixture','authority_url':'https://example.org/metadata'}}}
        (cache/'fixture.source.json').write_text(json.dumps(receipt))
        assert validate_case(case, {'CHEY_CPMG'}, cache)['table2_type_checks']
        receipt['identifier_aliases']['fixture-paper']['doi']='10.wrong'
        (cache/'fixture.source.json').write_text(json.dumps(receipt))
        try: validate_case(case, {'CHEY_CPMG'}, cache)
        except ValueError: pass
        else: raise AssertionError('Wrong primary publication accepted')
        receipt['identifier_aliases']['fixture-paper']['doi']='10.fixture'
        (cache/'fixture.source.json').write_text(json.dumps(receipt))
        for change in ({'residues_author': 'D1, C2'}, {'locator': 'Results'}, {'access': 'abstract_only', 'temperature': '298 K'}):
            (case / 'rows.jsonl').write_text(json.dumps({**evidence[0], **change}) + '\n')
            try: validate_case(case, {'CHEY_CPMG'}, cache)
            except ValueError: pass
            else: raise AssertionError('Missing source/type safeguard')
        assert len(paper_list([{'doi': str(i), 'title': str(i), 'row_id': 'r' + str(i)} for i in range(10)])[:8]) == 8
    print('Exporter workbook/HTML/type-check safeguards PASS; synthetic fixtures only.')


def main():
    try: import openpyxl
    except ModuleNotFoundError:
        if RUNTIME.is_file() and Path(sys.executable).resolve() != RUNTIME.resolve():
            os.execv(str(RUNTIME), [str(RUNTIME), __file__, *sys.argv[1:]])
        raise RuntimeError('Installed workbook runtime unavailable; do not install a dependency')
    parser = argparse.ArgumentParser(); parser.add_argument('--self-test', action='store_true')
    parser.add_argument('--verify-export', action='store_true')
    parser.add_argument('--output', type=Path); parser.add_argument('--cases', nargs='+', type=Path)
    args = parser.parse_args()
    if (EVAL / 'STOP').exists(): raise ValueError('STOP exists')
    if args.self_test: return self_test()
    if not args.cases or (not args.verify_export and not args.output): parser.error('--output and --cases required')
    output = args.output.resolve() if args.output else None
    if not args.verify_export and (output.parent != (WORK / 'outputs').resolve() or
            output.name not in ('gina_review_20261005', 'atlas_10_20261008')):
        raise ValueError('Only registered package destinations allowed')
    dev_ids = set((WORK / 'config/dev_ids.txt').read_text().split()); loaded = []
    for case in args.cases:
        case = case.resolve()
        entry = read(case / 'input.json')['entry_id']
        permitted = entry in dev_ids or (args.verify_export and (case / '.released').exists())
        if case.parent != (WORK / 'cases').resolve() or not permitted:
            raise ValueError('Only development or officially released case paths allowed')
        score = subprocess.run([sys.executable, str(EVAL / 'eval/score.py'), str(case)], capture_output=True, text=True)
        if score.returncode: raise ValueError('Current source integrity check failed')
        loaded.append(validate_case(case, dev_ids | {entry} if args.verify_export else dev_ids, WORK / 'sources_cache'))
    if args.verify_export:
        with tempfile.TemporaryDirectory(dir=WORK / 'ledger/tests') as temp:
            render(Path(temp), loaded)
            book = openpyxl.load_workbook(Path(temp) / 'atlas.xlsx')
            assert book['evidence_rows'].max_row == 1 + sum(len(c['rows']) for c in loaded)
        print(json.dumps({'reexport_verified': True, 'cases': len(loaded)})); return
    expected = {'AQADK_CPMG', 'CHEY_CPMG'} if output.name == 'gina_review_20261005' else dev_ids
    if len(loaded) != len(expected) or {c['entry'] for c in loaded} != expected:
        raise ValueError('Package entry set incomplete or duplicated')
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output.parent) as temp:
        staged = Path(temp) / output.name; render(staged, loaded)
        if (EVAL / 'STOP').exists(): raise ValueError('STOP before package promotion')
        if output.exists():
            backup = WORK / 'ledger/export_backups' / (datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S%f') + '_' + output.name)
            backup.parent.mkdir(parents=True, exist_ok=True); output.rename(backup)
        staged.rename(output)
    print(json.dumps({'exported': str(output), 'entries': sorted(expected), 'scientific_acceptance': 'pending_Gina'}))


if __name__ == '__main__': main()
